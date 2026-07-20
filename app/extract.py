"""extract.py — Text aus hochgeladenen Dokumenten (PDF, Word, Excel, PowerPoint).

Bewusst SERVERSEITIGE Text-Extraktion statt nativer Datei-Uploads ans Modell:
das Gespräch schickt bei jeder Runde den ganzen Verlauf mit — als Text bleiben
Kosten und Grösse berechenbar, und der Verlauf ist les- und prüfbar.

Grenze (ehrlich): Layout, Diagramme und Bilder gehen verloren; ein reiner Scan
ohne Textebene liefert nichts. Das melden wir zurück, statt es zu verschweigen.
"""
from __future__ import annotations

import io

MAX_BYTES = 10 * 1024 * 1024      # 10 MB je Datei
MAX_CHARS = 40_000                # ~10k Tokens je Dokument
KINDS = {".pdf": "PDF", ".docx": "Word", ".xlsx": "Excel", ".pptx": "PowerPoint"}


def kind_of(filename: str) -> str:
    """Dateityp aus der Endung; '' wenn nicht unterstützt."""
    low = (filename or "").lower()
    for ext, name in KINDS.items():
        if low.endswith(ext):
            return name
    return ""


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    return "\n".join((p.extract_text() or "") for p in r.pages)


def _docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:                                  # Tabellen zeilenweise
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Blatt: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def _pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"## Folie {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
    return "\n".join(out)


_READERS = {"PDF": _pdf, "Word": _docx, "Excel": _xlsx, "PowerPoint": _pptx}


def extract(filename: str, data: bytes):
    """-> (kind, text, hinweis). text == '' bedeutet: nichts Verwertbares."""
    kind = kind_of(filename)
    if not kind:
        return "", "", "Dateityp nicht unterstützt (nur PDF, Word, Excel, PowerPoint)."
    if len(data) > MAX_BYTES:
        return kind, "", f"Datei zu gross (max. {MAX_BYTES // 1024 // 1024} MB)."
    try:
        text = (_READERS[kind](data) or "").strip()
    except ImportError:
        return kind, "", f"{kind}-Unterstützung ist auf dem Server nicht installiert."
    except Exception as e:
        return kind, "", f"Datei konnte nicht gelesen werden ({type(e).__name__})."
    if not text:
        return kind, "", ("Kein Text gefunden — vermutlich ein Scan oder ein rein "
                          "grafisches Dokument. Bilder werden nicht ausgewertet.")
    note = ""
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
        note = f"gekürzt auf {MAX_CHARS} Zeichen"
    return kind, text, note
